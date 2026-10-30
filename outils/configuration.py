#!/usr/bin/env python3
# mon-ecran-version: 6df8f5c1
# Lit le classeur de configuration (configuration.xlsx) et écrit les réglages du dock dans config/ :
#   config/sport.json, actus.json, marches.json, videos.json, tv.json, cotes.json, accueil.json, apparence.json, applis.json
#   config/config.js                     → les mêmes réglages en un seul fichier, chargé par les pages (window.MON_ECRAN_CONFIG)
#   config/chaines-tv.xml                → liste des chaînes pour le programme TV (lue par la tâche « Programme TV sport »)
#   config/rapport.json                  → résumé + résultat de chaque test (affiché dans la page État)
#   config/configuration-resultat.xlsx   → copie du classeur avec la colonne « Résultat » remplie
# Usage : python3 outils/configuration.py configuration.xlsx [--sortie config] [--tester]
#   --tester : essaie réellement chaque nouveauté cochée et chaque ajout (flux, ESPN, Yahoo, YouTube, météo…)
import sys, os, re, json, unicodedata, datetime, urllib.request, urllib.parse, urllib.error
from openpyxl import load_workbook

args = sys.argv[1:]
POS = [a for i, a in enumerate(args) if not a.startswith('--') and (i == 0 or args[i - 1] != '--sortie')]
SRC = POS[0] if POS else 'configuration.xlsx'
OUT = args[args.index('--sortie') + 1] if '--sortie' in args else 'config'
TESTER = '--tester' in args
UA = 'Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/605.1.15 (KHTML, like Gecko) Version/17.0 Safari/605.1.15'

def slug(s):
    s = unicodedata.normalize('NFD', str(s or '').lower()).encode('ascii', 'ignore').decode()
    return re.sub(r'[^a-z0-9]+', '-', s).strip('-')
def txt(v): return str(v).strip() if v is not None else ''
def oui(v): return txt(v).lower() in ('oui', 'o', 'x', 'yes', '1', '✓', 'vrai', 'true')
def non(v): return txt(v).lower() in ('non', 'n', 'no', '0', 'faux', 'false')

# ---------- Lecture des onglets ----------
def lire(wb):
    mods = {}
    for ws in wb.worksheets:
        a1 = txt(ws['A1'].value)
        if not a1.startswith('module='): continue
        mid = a1.split(';')[0][7:]
        H = {txt(c.value): c.column for c in ws[5] if c.value}
        c2 = next((H[k] for k in ('Favori', 'Widget accueil', 'Abonné', 'Par défaut') if k in H), None)
        m = dict(id=mid, ws=ws, H=H, sections=[], ajouts=[])
        sec = None
        for r in range(6, ws.max_row + 1):
            i = txt(ws.cell(r, 1).value)
            if not i or i == 'exemple': continue
            if i.startswith('sec:'):
                sec = dict(id=i.split(':', 2)[2], oui=txt(ws.cell(r, H['Oui']).value), tech=json.loads(ws.cell(r, H['Technique']).value), rows=[], r=r)
                m['sections'].append(sec); continue
            if i.startswith('ajout:'):
                v = [txt(ws.cell(r, c).value) for c in range(2, 5)]
                if v[1] or v[2]:
                    m['ajouts'].append(dict(r=r, section=v[0], nom=v[1], source=v[2] if H.get('Oui') != 4 else '', oui=txt(ws.cell(r, H['Oui']).value)))
                continue
            t = json.loads(ws.cell(r, H['Technique']).value)
            sec['rows'].append(dict(id=i, r=r, oui=txt(ws.cell(r, H['Oui']).value), oui2=txt(ws.cell(r, c2).value) if c2 else '',
                                    cells={h: txt(ws.cell(r, c).value) for h, c in H.items()},
                                    support=t.get('support', 'ok'), d=t.get('def'), nom=txt(ws.cell(r, 3).value),
                                    resultat=txt(ws.cell(r, H['Résultat']).value)))
        mods[mid] = m
    return mods

def choisi(sec, row): return oui(row['oui']) or (oui(sec['oui']) and not non(row['oui']))

RAPPORT = dict(date=datetime.datetime.now(datetime.timezone.utc).isoformat(timespec='seconds'), source=os.path.basename(SRC), modules={}, tests=[], erreurs=[])
RES = {}          # (feuille, ligne) → texte du résultat
def note(m, r, statut, msg, nom=''):
    sym = {'ok': '✅', 'warn': '⚠️', 'ko': '❌', 'info': 'ℹ️'}[statut]
    RES[(m['ws'].title, r)] = f'{sym} {msg}'
    RAPPORT['tests'].append(dict(onglet=m['ws'].title, ligne=r, nom=nom, statut=statut, message=msg))

def selection(m, a_tester=None):
    """Sections choisies (dans l'ordre du classeur) avec leurs éléments ; les éléments « à développer » / « non disponibles » sont signalés et écartés."""
    out = []
    for s in m['sections']:
        rows = []
        for x in s['rows']:
            if not choisi(s, x): continue
            if x['support'] == 'non': note(m, x['r'], 'ko', 'non disponible : pas de source gratuite fiable', x['nom']); continue
            if x['support'] == 'dev': note(m, x['r'], 'info', "données disponibles, affichage à développer dans l'appli", x['nom']); continue
            rows.append(x)
            if a_tester and x['resultat'].startswith('à vérifier'): a_tester(m, x)
        out.append((s, rows))
    return out

# ---------- Tests réseau (seulement avec --tester) ----------
def get(url, accept='*/*', timeout=15):
    req = urllib.request.Request(url, headers={'user-agent': UA, 'accept': accept})
    with urllib.request.urlopen(req, timeout=timeout) as r: return r.read().decode('utf-8', 'replace'), r.geturl()
def essai(fn):
    def w(*a):
        if not TESTER: return None
        try: return fn(*a)
        except Exception as e: return ('warn', f'non testé : site injoignable pour l\'instant ({type(e).__name__}) — gardé, nouvel essai au prochain dépôt')
    return w
@essai
def t_flux(url, site=None):
    body, final = get(url, 'application/rss+xml, application/atom+xml, application/xml, text/xml, */*')
    n = len(re.findall(r'<(item|entry)\b', body))
    if n: return ('ok', f'{n} articles dans le flux')
    for cand in [site, final]:
        if not cand: continue
        html = body if cand == final else get(cand)[0]
        m = re.search(r'<link[^>]+type=["\']application/(?:rss|atom)\+xml["\'][^>]*href=["\']([^"\']+)', html, re.I) or \
            re.search(r'<link[^>]+href=["\']([^"\']+)["\'][^>]*type=["\']application/(?:rss|atom)\+xml', html, re.I)
        if m:
            return ('warn', f'flux trouvé sur la page : {urllib.parse.urljoin(cand, m.group(1))}', urllib.parse.urljoin(cand, m.group(1)))
    return ('ko', "aucun article ni flux RSS trouvé à cette adresse")
SPORTS_ESPN = ['soccer', 'rugby', 'basketball', 'football', 'hockey', 'baseball', 'racing', 'tennis', 'golf', 'mma']
@essai
def t_espn(cle, sport=None):
    joint = False
    for sp in ([sport] if sport else SPORTS_ESPN):
        try:
            body, _ = get(f'https://site.api.espn.com/apis/site/v2/sports/{sp}/{cle}/scoreboard', 'application/json')
            joint = True; j = json.loads(body)
            if j.get('leagues') or j.get('events') is not None:
                nom = (j.get('leagues') or [{}])[0].get('name', '')
                return ('ok', f'ESPN répond ({sp}) {nom}'.strip(), sp)
        except urllib.error.HTTPError: joint = True; continue
        except Exception: continue
    if not joint: return ('warn', 'non testé : ESPN injoignable pour l\'instant — gardé, nouvel essai au prochain dépôt')
    return ('ko', f'clé ESPN « {cle} » inconnue')
@essai
def t_yahoo(code):
    body, _ = get(f'https://query1.finance.yahoo.com/v8/finance/chart/{urllib.parse.quote(code)}?range=5d&interval=1d', 'application/json')
    j = json.loads(body); res = (j.get('chart') or {}).get('result')
    if res: return ('ok', f"cours trouvé ({res[0]['meta'].get('currency', '')} {res[0]['meta'].get('regularMarketPrice', '')})")
    return ('ko', f'code « {code} » inconnu de Yahoo Finance')
@essai
def t_isin(isin):
    body, _ = get(f'https://query2.finance.yahoo.com/v1/finance/search?q={isin}&quotesCount=5&newsCount=0', 'application/json')
    q = json.loads(body).get('quotes') or []
    if q: return ('ok', f"reconnu : {q[0].get('symbol')} · {q[0].get('shortname') or q[0].get('longname') or ''}", q[0].get('symbol'))
    return ('warn', 'ISIN non trouvé : ajoute un code de cotation dans la colonne Source')
@essai
def t_youtube(ident):
    if re.fullmatch(r'UC[\w-]{22}', ident): uc = ident
    else:
        url = ident if ident.startswith('http') else 'https://www.youtube.com/' + (ident if ident.startswith('@') else '@' + ident)
        html, _ = get(url)
        m = re.search(r'"(?:channelId|externalId)":"(UC[\w-]{22})"', html) or re.search(r'channel_id=(UC[\w-]{22})', html)
        if not m: return ('ko', 'chaîne introuvable')
        uc = m.group(1)
    body, _ = get(f'https://www.youtube.com/feeds/videos.xml?channel_id={uc}')
    n = len(re.findall(r'<entry>', body))
    return ('ok', f'chaîne {uc} · {n} vidéos récentes', uc) if n else ('warn', f'chaîne {uc} trouvée, aucune vidéo récente', uc)
@essai
def t_ville(nom, region=''):
    body, _ = get('https://geocoding-api.open-meteo.com/v1/search?' + urllib.parse.urlencode({'name': nom, 'count': 5, 'language': 'fr'}), 'application/json')
    res = json.loads(body).get('results') or []
    if region: res = sorted(res, key=lambda x: slug(region) not in slug(' '.join(str(x.get(k, '')) for k in ('admin1', 'admin2'))))
    if not res: return ('ko', 'ville introuvable')
    v = res[0]; return ('ok', f"{v['name']} · {v.get('admin2') or v.get('admin1', '')} ({v['latitude']:.4f}, {v['longitude']:.4f})", v)
@essai
def t_tmdb(nom):
    tok = os.environ.get('TMDB_TOKEN', '').strip()
    if not tok: return ('warn', 'non testé : secret TMDB_TOKEN absent')
    req = urllib.request.Request('https://api.themoviedb.org/3/watch/providers/movie?watch_region=FR&language=fr-FR', headers={'authorization': f'Bearer {tok}', 'accept': 'application/json'})
    with urllib.request.urlopen(req, timeout=15) as r: ps = json.loads(r.read()).get('results', [])
    hit = [p['provider_name'] for p in ps if slug(nom) in slug(p['provider_name'])]
    return ('ok', f'connue de TMDB en France : {hit[0]}') if hit else ('ko', 'plateforme absente de TMDB pour la France')
@essai
def t_odds(cle):
    key = os.environ.get('ODDS_API_KEY', '').strip()
    if not key: return ('warn', 'non testé : secret ODDS_API_KEY absent')
    body, _ = get(f'https://api.the-odds-api.com/v4/sports?apiKey={key}&all=true', 'application/json')   # gratuit, ne consomme pas de crédit
    ks = {s['key']: s for s in json.loads(body)}
    if cle in ks: return ('ok', f"proposée par l'API : {ks[cle]['title']}" + ('' if ks[cle].get('active') else ' (hors saison)'))
    return ('ko', "clé absente de l'API")

def appliquer_test(m, r, nom, res):
    if res: note(m, r, res[0], res[1], nom)
    return res

# ---------- Modules ----------
def section_pour(m, nom):
    """Section d'un ajout : existante (par libellé ou id), sinon nouvelle."""
    k = slug(nom)
    for s in m['sections']:
        if k in (slug(s['tech'].get('label', '')), slug(s['id']), slug(s['tech'].get('label', '')).split('-page-')[0]): return s
    s = dict(id=k or 'ajouts', oui='', tech={'id': k or 'ajouts', 'label': nom or 'Mes ajouts'}, rows=[], r=None, nouvelle=True)
    m['sections'].append(s); return s

def ajouts(m):
    for a in m['ajouts']:
        if non(a['oui']): continue
        if not oui(a['oui']): note(m, a['r'], 'info', 'mets « Oui » pour l\'ajouter', a['nom']); continue
        yield a

def m_sport(M):
    m = M['sport']
    def tst(m, x): appliquer_test(m, x['r'], x['nom'], t_espn(x['d']['id'], x['d']['sport']))
    for a in ajouts(m):
        s = section_pour(m, a['section'] or 'Mes ajouts'); d = {'id': a['source'] or slug(a['nom']), 'region': s['tech']['id'], 'flag': '🏳️', 'sport': 'soccer', 'group': 'foot', 'name': a['nom'], 'ajout': True}
        if not a['source']: note(m, a['r'], 'warn', 'indique la clé ESPN dans la colonne Source (ex. sco.1)', a['nom']); continue
        res = t_espn(a['source'])
        if res and res[0] == 'ok': d['sport'] = res[2]
        appliquer_test(m, a['r'], a['nom'], res)
        if res and res[0] == 'ko': continue
        s['rows'].append(dict(id='ajout', r=a['r'], oui='Oui', oui2='', support='ok', d=d, nom=a['nom'], resultat=''))
    sel = selection(m, tst)
    regions = [s['tech'] for s, rows in sel if rows]
    leagues = [x['d'] for s, rows in sel for x in rows]
    return {'regions': regions, 'leagues': leagues}     # favoris : choisis dans l'appli (étoiles), partagés par la synchronisation

def m_actus(m):
    def tst(m, x):
        res = t_flux(x['d']['url'], x['d'].get('site'))
        appliquer_test(m, x['r'], x['nom'], res)
    for a in ajouts(m):
        s = section_pour(m, a['section'] or 'Mes ajouts'); src = a['source']
        if src.startswith('http'):
            d = {'name': a['nom'] or src, 'color': '#6b7280', 'url': src, 'site': src}
            res = t_flux(src, src)
            if res and res[0] == 'warn' and len(res) > 2: d['url'] = res[2]
            appliquer_test(m, a['r'], a['nom'], res)
            if res and res[0] == 'ko': continue
        else:
            q = src or a['nom']; hl = 'hl=en-US&gl=US&ceid=US:en' if m['id'] == 'actus-en' else 'hl=fr&gl=FR&ceid=FR:fr'
            d = {'name': a['nom'] or q, 'color': '#6b7280', 'url': f'https://news.google.com/rss/search?q={urllib.parse.quote(q)}&{hl}'}
            appliquer_test(m, a['r'], a['nom'], t_flux(d['url']) or ('info', 'flux Google Actualités sur le mot-clé « ' + q + ' »'))
        d['ajout'] = True
        s['rows'].append(dict(id='ajout', r=a['r'], oui='Oui', oui2='', support='ok', d=d, nom=a['nom'], resultat=''))
    out = {'principal': [], 'autres': []}
    for s, rows in selection(m, tst):
        if not rows: continue
        out[s['tech'].get('liste', 'principal')].append({'id': s['tech']['id'], 'label': s['tech']['label'], 'feeds': [x['d'] for x in rows]})
    return out

def m_marches(m):
    def tst(m, x):
        d = x['d']; appliquer_test(m, x['r'], x['nom'], t_yahoo(d['code']) if d.get('code') else t_isin(d['isin']))
    for a in ajouts(m):
        src = a['source'].upper().replace(' ', '')
        if re.fullmatch(r'[A-Z]{2}[A-Z0-9]{9}\d', src):
            res = t_isin(src); d = {'isin': src, 'nom': a['nom'] or src}
            if res and len(res) > 2 and res[2]: d['code'] = res[2]
            s = next(s for s in m['sections'] if s['tech'].get('liste') == 'suivi')
        else:
            res = t_yahoo(a['source']); d = {'id': slug(a['nom'] or a['source']), 'code': a['source'], 'nom': a['nom'] or a['source']}
            s = section_pour(m, a['section'] or 'Mes ajouts'); s['tech'].setdefault('liste', 'indices')
        appliquer_test(m, a['r'], a['nom'], res)
        if res and res[0] == 'ko': continue
        s['rows'].append(dict(id='ajout', r=a['r'], oui='Oui', oui2='', support='ok', d=d, nom=a['nom'], resultat=''))
    indices, suivi, widget = [], [], []
    for s, rows in selection(m, tst):
        for x in rows:
            d = dict(x['d'])
            if s['tech'].get('liste') == 'suivi': suivi.append(d); continue
            if oui(x['oui2']): d['accueil'] = True; widget.append(d['code'])
            indices.append(d)
    return {'indices': indices, 'suivi': suivi}, widget

def m_videos(M):
    m = M['videos']
    def tst(m, x): appliquer_test(m, x['r'], x['nom'], t_tmdb(x['d']['nom']))
    for a in ajouts(m):
        nom = a['nom']; d = {'id': slug(nom), 'nom': nom, 'court': nom[:3], 'coul': '#6b7280', 'tmdb': '^' + re.escape((a['source'] or nom).lower()), 'web': ''}
        res = t_tmdb(a['source'] or nom); appliquer_test(m, a['r'], nom, res)
        if res and res[0] == 'ko': continue
        m['sections'][0]['rows'].append(dict(id='ajout', r=a['r'], oui='Oui', oui2='', support='ok', d=d, nom=nom, resultat=''))
    pl = []
    for s, rows in selection(m, tst):
        for x in rows:
            pl.append(dict(x['d']))        # « abo » : valeur de départ ; tes abonnements se cochent dans la page Vidéos
    y = M['youtube']
    def tsty(m, x):
        res = t_youtube(x['d']['id']); appliquer_test(m, x['r'], x['nom'], res)
        if res and res[0] in ('ok', 'warn') and len(res) > 2: x['d'] = {'id': res[2], 'nom': x['d']['nom']}
    for a in ajouts(y):
        section_pour(y, a['section'] or 'Mes ajouts')['rows'].append(dict(id='ajout', r=a['r'], oui='Oui', oui2='', support='ok',
                                                                         d={'id': a['source'] or a['nom'], 'nom': a['nom'] or a['source']}, nom=a['nom'], resultat='à vérifier'))
    ch = [x['d'] for s, rows in selection(y, tsty) for x in rows]
    return {'plateformes': pl, 'chaines_yt': ch}

def m_tv(m):
    for a in ajouts(m):
        section_pour(m, a['section'] or 'Chaînes gratuites et autres')['rows'].append(dict(id='ajout', r=a['r'], oui='Oui', oui2='', support='ok', d={'nom': a['nom'], 'sources': []}, nom=a['nom'], resultat='à vérifier'))
    src, sans = [], []
    for s, rows in selection(m):
        for x in rows:
            if x['d']['sources']: src += x['d']['sources']
            else:
                sans.append(x['d']['nom']); note(m, x['r'], 'warn', 'chaîne retenue : sa source de programme sera cherchée dans les guides TV publics', x['nom'])
    src = [{k: v for k, v in c.items() if k != 'ordre'} for c in sorted(src, key=lambda c: c.get('ordre', 1e9))]
    return {'chaines': src, 'a_trouver': sans}

def m_cotes(m):
    def tst(m, x):
        if isinstance(x['d'], dict): appliquer_test(m, x['r'], x['nom'], t_odds(x['d']['cle']))
        else: note(m, x['r'], 'warn', "site à confirmer : il apparaîtra s'il fait partie de la région « fr » de l'API", x['nom'])
    for a in ajouts(m):
        s = section_pour(m, a['section'] or 'Compétitions relevées')
        if s['id'] == 'sites': d = a['source'] or slug(a['nom'])
        else: d = {'id': a['source'] or slug(a['nom']), 'cle': a['source'], 're': {'re': re.escape(a['nom']), 'flags': 'i'}}
        s['rows'].append(dict(id='ajout', r=a['r'], oui='Oui', oui2='', support='ok', d=d, nom=a['nom'], resultat='à vérifier'))
    sites, comps, seuil, ref = [], {}, [], []
    for s, rows in selection(m, tst):
        for x in rows:
            if s['id'] == 'sites': sites.append(x['d'])
            elif s['id'] == 'rapport_seuil': seuil.append(x['d'])
            elif s['id'] == 'rapport_ref': ref.append(x['d'])
            else: d = dict(x['d']); comps[d.pop('id')] = d
    for sid, l in (('rapport_seuil', seuil), ('rapport_ref', ref)):
        if len(l) > 1:
            x = next(x for s, rows in selection(m) if s['id'] == sid for x in rows)
            note(m, x['r'], 'warn', 'un seul choix possible : le premier coché est gardé', x['nom'])
    return {'sites': sites, 'comps': comps, 'rapport': {'seuil': (seuil or [0.05])[0], 'reference': (ref or ['moyenne'])[0]}}

def m_accueil(m):
    S = {s['id']: (s, rows) for s, rows in selection(m)}
    une = [x['d'] for x in S['une'][1]]
    if len(une) > 1: note(m, S['une'][1][1]['r'], 'warn', 'une seule source « À la une » : la première cochée est gardée', S['une'][1][1]['nom'])
    return {'pages': [x['d'] for x in S['pages'][1]], 'widgets': [x['d'] for x in S['widgets'][1]], 'une': une[0] if une else None}

def m_applis(m):
    for a in ajouts(m):
        s = section_pour(m, a['section'] or 'Mes ajouts')
        if s.get('nouvelle'): s['tech'] = {'id': s['id'], 'label': a['section'] or 'Mes ajouts', 'type': 'dossier'}
        if not a['source']: note(m, a['r'], 'warn', "indique au moins le lien iPhone (ex. spotify://) ou un lien web", a['nom']); continue
        d = {'label': a['nom'], 'mono': a['nom'][:2], 'color': '--t-slate'}
        d['app' if not a['source'].startswith('http') else 'web'] = a['source']
        note(m, a['r'], 'ok', 'ajoutée (teste-la depuis le dock)', a['nom'])
        s['rows'].append(dict(id='ajout', r=a['r'], oui='Oui', oui2='', support='ok', d=d, nom=a['nom'], resultat='', cells={}))
    out = {'accueil': [], 'dossiers': {}, 'bas_gauche': [], 'bas_droite': [], 'navigation': []}
    for s, rows in selection(m):
        items = []
        for x in rows:
            d = dict(x['d']); c = x.get('cells') or {}
            if not d.get('type') and not d.get('action') and c:          # liens modifiables dans le tableau
                for k, h in (('app', 'Lien iPhone'), ('android', 'Lien Android'), ('web', 'Lien web')):
                    v = c.get(h, '')
                    if v: d[k] = v
                    else: d.pop(k, None)
            items.append(d)
        t = s['tech']
        if t.get('type') == 'dossier':
            if items: out['dossiers'][t['id']] = {'title': t['label'], 'items': items}
        else: out[t['id']] = items
    for k in out['dossiers']:                                            # nouveau dossier : sa tuile sur l'accueil
        if not any(a.get('target') == k for a in out['accueil']): out['accueil'].append({'label': out['dossiers'][k]['title'], 'type': 'folder', 'target': k})
    out['accueil'] = [a for a in out['accueil'] if a.get('type') != 'folder' or a.get('target') in out['dossiers']]
    return out

def m_apparence(m):
    S = {s['id']: (s, rows) for s, rows in selection(m)}
    pal = [x['d'] for x in S['themes'][1]]
    defs = [x for s in m['sections'] if s['id'] == 'themes' for x in s['rows'] if oui(x['oui2'])]
    defaut = [x['d'] for x in defs]
    modes = [x['d'] for x in S['mode'][1]]
    if len(defs) > 1: note(m, defs[0]['r'], 'warn', f"{len(defs)} thèmes « Par défaut » : seul le premier ({defs[0]['nom']}) est gardé", defs[0]['nom'])
    if len(S['mode'][1]) > 1: note(m, S['mode'][1][0]['r'], 'warn', 'plusieurs modes cochés : seul le premier est gardé', S['mode'][1][0]['nom'])
    return {'palettes': pal, 'palette': (defaut or pal or ['bleu'])[0], 'mode': (modes or ['auto'])[0]}

# ---------- Programme principal ----------
def main():
    wb = load_workbook(SRC)
    M = lire(wb)
    manquants = [k for k in ('sport', 'actus', 'actus-en', 'marches', 'videos', 'youtube', 'tv', 'cotes', 'accueil', 'apparence', 'applis') if k not in M]
    if manquants: raise SystemExit(f'Onglets introuvables dans {SRC} : {", ".join(manquants)} (as-tu supprimé la colonne A ou renommé une ligne 1 ?)')
    fr, en = m_actus(M['actus']), m_actus(M['actus-en'])
    marches, widget = m_marches(M['marches'])
    accueil = m_accueil(M['accueil']); accueil['indices_widget'] = widget
    cfg = {
        'sport': m_sport(M),
        'actus': {'fr': fr['principal'], 'autres': fr['autres'], 'en': en['principal']},
        'marches': marches,
        'videos': m_videos(M),
        'tv': m_tv(M['tv']),
        'cotes': m_cotes(M['cotes']),
        'accueil': accueil,
        'apparence': m_apparence(M['apparence']),
        'applis': m_applis(M['applis']),
    }
    os.makedirs(OUT, exist_ok=True)
    for k, v in cfg.items():
        json.dump(v, open(os.path.join(OUT, f'{k}.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    # un seul fichier pour les pages (chargé avant leur code ; en cas d'absence, elles gardent leurs réglages intégrés)
    with open(os.path.join(OUT, 'config.js'), 'w', encoding='utf-8') as f:
        f.write('// Réglages du dock, écrits par la tâche « Configuration » à partir de configuration.xlsx — ne pas modifier à la main.\n')
        f.write('window.MON_ECRAN_CONFIG = ' + json.dumps(dict(cfg, genere=RAPPORT['date']), ensure_ascii=False, separators=(',', ':')) + ';\n')
    # liste des chaînes pour le programme TV (même format que outils/chaines-tv.xml)
    esc = lambda v: str(v).replace('&', '&amp;').replace('<', '&lt;').replace('"', '&quot;')
    with open(os.path.join(OUT, 'chaines-tv.xml'), 'w', encoding='utf-8') as f:
        f.write('<?xml version="1.0" encoding="UTF-8"?>\n<!-- écrit par la tâche « Configuration » (onglet Programme TV du classeur) -->\n<channels>\n')
        for c in cfg['tv']['chaines']:
            f.write(f'  <channel site="{esc(c["site"])}" site_id="{esc(c["site_id"])}" lang="{esc(c["lang"])}" xmltv_id="{esc(c["xmltv_id"])}">{esc(c["nom"])}</channel>\n')
        f.write('</channels>\n')
    for mid, m in M.items():
        RAPPORT['modules'][mid] = {'onglet': m['ws'].title, 'coches': sum(1 for s in m['sections'] for x in s['rows'] if choisi(s, x)), 'ajouts': len(m['ajouts'])}
    st = [t['statut'] for t in RAPPORT['tests']]
    RAPPORT['resume'] = {k: st.count(k) for k in ('ok', 'warn', 'ko', 'info')}
    RAPPORT['teste'] = TESTER
    json.dump(RAPPORT, open(os.path.join(OUT, 'rapport.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    # Copie du classeur avec les résultats
    for (feuille, r), v in RES.items():
        ws = wb[feuille]; c = next(c.column for c in ws[5] if txt(c.value) == 'Résultat'); ws.cell(r, c, v)
    wb.save(os.path.join(OUT, 'configuration-resultat.xlsx'))
    print(f"Configuration écrite dans {OUT}/ · {sum(m['coches'] for m in RAPPORT['modules'].values())} éléments cochés · tests : {RAPPORT['resume']}")

if __name__ == '__main__':
    main()
