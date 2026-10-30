# mon-ecran-version: 7302bf3b
# Construit le classeur de configuration « Mon Écran » (openpyxl).
# Usage : python3 construire.py reference.json sortie.xlsx [ancien.xlsx pour reprendre les choix]
import sys, json
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter as L
from openpyxl.comments import Comment
from modele import construire, VERSION_CLASSEUR
from tuto import LISEZMOI, TUTO_REMPLIR, TUTO_PRIVE, JOURNAL

NAVY = '274C77'; NAVY2 = '1B3656'; SEC = 'DCE6F1'; INPUT = 'FFF4C2'; OK = 'E2F0D9'; GRIS = '6B7280'; LIGNE = 'C9D3DD'
F = lambda **k: Font(name='Arial', **{'size': 10, **k})
FILL = lambda c: PatternFill('solid', start_color=c, end_color=c)
fin = Side(style='thin', color=LIGNE)
BORD = Border(bottom=fin)
WR = Alignment(wrap_text=True, vertical='top')
MID = Alignment(vertical='center')
CENTER = Alignment(horizontal='center', vertical='center')

LIENS = ('Lien iPhone', 'Lien Android', 'Lien web')       # modifiables dans l'appli : repris tels quels

def anciens_choix(path):
    """Reprend les « Oui / Non », la 2e case, les liens d'applis et les ajouts d'un ancien classeur, par identifiant
    (y compris ceux reportés depuis l'appli par la tâche Configuration)."""
    if not path: return {}, {}
    wb = load_workbook(path); choix = {}; ajouts = {}
    for ws in wb.worksheets:
        a1 = str(ws['A1'].value or '')
        if not a1.startswith('module='): continue
        mod = a1.split(';')[0][7:]
        hdr = {str(c.value).strip(): c.column for c in ws[5] if c.value}
        co, c2 = hdr.get('Oui'), next((v for k, v in hdr.items() if k in ('Favori', 'Widget accueil', 'Abonné', 'Par défaut')), None)
        for r in range(6, ws.max_row + 1):
            i = ws.cell(r, 1).value
            if not i or i == 'exemple': continue
            if str(i).startswith('ajout:'):
                vals = [ws.cell(r, c).value for c in range(2, 6)]
                if any(vals[:2]): ajouts.setdefault(mod, []).append([ws.cell(r, c).value for c in range(2, ws.max_column + 1)])
                continue
            liens = {h: ws.cell(r, c).value for h, c in hdr.items() if h in LIENS}
            choix[i] = (ws.cell(r, co).value if co else None, ws.cell(r, c2).value if c2 else None, liens)
    return choix, ajouts

def feuille_module(wb, m, choix, ajouts_prec):
    ws = wb.create_sheet(m['feuille'])
    cols = m['cols']; c2 = m.get('coche2')
    heads = ['ID', 'Section'] + [h for _, h, _ in cols] + ['Oui'] + ([c2] if c2 else []) + ['Résultat', 'Note', 'Technique']
    widths = [14, 26] + [w for _, _, w in cols] + [8] + ([12] if c2 else []) + [26, 34, 10]
    cOui = 3 + len(cols); c2c = cOui + 1 if c2 else None; cRes = cOui + (2 if c2 else 1); cNote = cRes + 1; cTech = cNote + 1
    last = cNote
    for i, w in enumerate(widths, 1): ws.column_dimensions[L(i)].width = w
    ws.column_dimensions['A'].hidden = True; ws.column_dimensions[L(cTech)].hidden = True
    ws['A1'] = f"module={m['id']};version={VERSION_CLASSEUR}"
    ws.merge_cells(start_row=1, start_column=2, end_row=1, end_column=last)
    t = ws.cell(1, 2, m['titre']); t.font = F(size=15, bold=True, color='FFFFFF'); t.fill = FILL(NAVY); t.alignment = MID
    for c in range(2, last + 1): ws.cell(1, c).fill = FILL(NAVY)
    ws.row_dimensions[1].height = 30
    ws.merge_cells(start_row=2, start_column=2, end_row=2, end_column=last)
    t = ws.cell(2, 2, m['intro']); t.font = F(color='1F2937'); t.alignment = WR
    ws.row_dimensions[2].height = 44
    ws.merge_cells(start_row=3, start_column=2, end_row=3, end_column=last)
    leg = ("À remplir : cases jaunes (« Oui » ou « Non »). Ligne « ▸ toute la section » = tous ses éléments, sauf ceux marqués « Non ». "
           "Un seul élément coché suffit pour que sa section apparaisse. Colonne « Résultat » : remplie par la tâche GitHub après dépôt du classeur.")
    t = ws.cell(3, 2, leg); t.font = F(size=9, italic=True, color=GRIS); t.alignment = WR
    ws.row_dimensions[3].height = 28
    for i, h in enumerate(heads, 1):
        c = ws.cell(5, i, h); c.font = F(bold=True, color='FFFFFF'); c.fill = FILL(NAVY2); c.alignment = CENTER if h in ('Oui', c2) else MID
    ws.row_dimensions[5].height = 22
    r = 6
    def coche(cell, v):
        cell.value = v or None; cell.fill = FILL(INPUT); cell.alignment = CENTER; cell.font = F(bold=True)
    for s in m['sections']:
        sid = f"sec:{m['id']}:{s['id']}"
        ws.cell(r, 1, sid); ws.cell(r, 2, s['label']).font = F(bold=True, color=NAVY2)
        ws.cell(r, 3, '▸ toute la section').font = F(italic=True, color=NAVY)
        for c in range(2, last + 1): ws.cell(r, c).fill = FILL(SEC); ws.cell(r, c).border = BORD
        prev = choix.get(sid, (None, None))
        coche(ws.cell(r, cOui), prev[0])
        ws.cell(r, cTech, json.dumps(s['tech'], ensure_ascii=False))
        if s.get('note'): ws.cell(r, cNote, s['note']).font = F(size=9, color=GRIS)
        r += 1
        for x in s['rows']:
            ws.cell(r, 1, x['id']); ws.cell(r, 2, s['label']).font = F(size=9, color='9CA3AF')
            p = choix.get(x['id'])
            for j, (k, h, _) in enumerate(cols):
                v = x['vals'].get(k, '')
                if p and h in p[2]: v = p[2][h] or ''
                c = ws.cell(r, 3 + j, v if v != '' else None); c.font = F(); c.alignment = MID
            coche(ws.cell(r, cOui), p[0] if p else x['oui'])
            if c2c: coche(ws.cell(r, c2c), p[1] if p else x['oui2'])
            res = ws.cell(r, cRes, x['resultat'] or None); res.font = F(size=9, color=GRIS)
            ws.cell(r, cNote, x['note'] or None).font = F(size=9, color=GRIS)
            ws.cell(r, cTech, json.dumps({'support': x['support'], 'def': x['tech']}, ensure_ascii=False))
            ws.row_dimensions[r].outline_level = 1
            for c in range(2, last + 1): ws.cell(r, c).border = BORD
            r += 1
    fin_table = r - 1
    ws.auto_filter.ref = f"B5:{L(last)}{fin_table}"
    # Mise en forme selon le choix
    rng = f"B6:{L(last)}{fin_table}"
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'${L(cOui)}6="Oui"'], fill=FILL(OK)))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'${L(cOui)}6="Non"'], font=Font(color='9CA3AF', strike=True)))
    # ---------- Mes ajouts ----------
    if m.get('ajout'):
        a_sec, a_nom, a_src, _, ex_nom, ex_src = m['ajout']
        r += 2
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=last)
        t = ws.cell(r, 2, "MES AJOUTS — tape ce que tu veux ajouter : la tâche essaie de le récupérer et écrit le résultat (✅ / ⚠️ / ❌) dans la colonne « Résultat »")
        t.font = F(bold=True, color='FFFFFF'); t.fill = FILL('8B5E00')
        for c in range(2, last + 1): ws.cell(r, c).fill = FILL('8B5E00')
        r += 1
        ah = {2: a_sec, 3: a_nom}
        if len(cols) >= 2: ah[4] = a_src
        ah[cOui] = 'Oui'; ah[cRes] = 'Résultat'; ah[cNote] = 'Note'
        for c, h in ah.items():
            x = ws.cell(r, c, h); x.font = F(bold=True, color='FFFFFF'); x.fill = FILL(NAVY2); x.alignment = WR
        ws.row_dimensions[r].height = 30
        r += 1
        ws.cell(r, 1, 'exemple')
        ex_sec = m['sections'][0]['label']
        for c, v in ((2, ex_sec), (3, ex_nom), (4 if len(cols) >= 2 else None, ex_src), (cOui, 'Oui'), (cNote, 'Ligne d\'exemple : ignorée par la tâche')):
            if c: ws.cell(r, c, v).font = F(italic=True, color='9CA3AF')
        r += 1
        prec = ajouts_prec.get(m['id'], [])
        a0 = r
        for k in range(1, max(21, len(prec) + 6)):           # 20 lignes libres, plus si l'appli en a déjà ajouté davantage
            ws.cell(r, 1, f'ajout:{k}')
            for c in [2, 3] + ([4] if len(cols) >= 2 else []):
                ws.cell(r, c).fill = FILL(INPUT); ws.cell(r, c).font = F(); ws.cell(r, c).border = BORD
            coche(ws.cell(r, cOui), None)
            if k <= len(prec):
                for j, v in enumerate(prec[k - 1]): ws.cell(r, 2 + j, v)
            r += 1
        dv_end = r - 1
    else:
        dv_end = fin_table
    dv = DataValidation(type='list', formula1='"Oui,Non"', allow_blank=True, showErrorMessage=True,
                        errorTitle='Choix', error='Réponds « Oui », « Non » ou laisse vide.')
    ws.add_data_validation(dv); dv.add(f"{L(cOui)}6:{L(cOui)}{dv_end}")
    if c2c: dv.add(f"{L(c2c)}6:{L(c2c)}{fin_table}")
    ws.freeze_panes = 'C6'
    ws.sheet_properties.outlinePr.summaryBelow = False
    ws.sheet_view.zoomScale = 100
    ws.sheet_properties.tabColor = NAVY
    return dict(feuille=m['feuille'], oui=L(cOui), c2=L(c2c) if c2c else None, c2nom=c2, debut=6, fin=fin_table, titre=m['titre'])

def texte(ws, r, c, v, **k):
    x = ws.cell(r, c, v); x.font = F(**k); x.alignment = WR; return x

def feuille_lisezmoi(wb, infos):
    ws = wb.create_sheet('Lisez-moi', 0)
    for col, w in zip('ABCDE', (3, 30, 16, 16, 60)): ws.column_dimensions[col].width = w
    ws.merge_cells('B1:E1'); t = ws['B1']; t.value = f'Mon Écran · classeur de configuration (version {VERSION_CLASSEUR})'
    t.font = F(size=17, bold=True, color='FFFFFF'); t.alignment = MID
    for c in 'BCDE': ws[f'{c}1'].fill = FILL(NAVY)
    ws.row_dimensions[1].height = 36
    r = 3
    for kind, txt in LISEZMOI:
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
        if kind == 'h': texte(ws, r, 2, txt, size=12, bold=True, color=NAVY2); ws.row_dimensions[r].height = 22
        else: texte(ws, r, 2, txt); ws.row_dimensions[r].height = max(16, 15 * (len(txt) // 105 + 1))
        r += 1
    r += 1
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5); texte(ws, r, 2, 'Légende', size=12, bold=True, color=NAVY2); r += 1
    for fill, lab in ((INPUT, 'Case à remplir : « Oui », « Non » ou vide'), (OK, 'Ligne choisie (se colore toute seule)'), (SEC, 'Ligne de section : « ▸ toute la section »'),
                      ('FFFFFF', 'Résultat : ✅ fonctionne · ⚠️ trouvé, à vérifier · ❌ introuvable · « à vérifier » = pas encore testé')):
        ws.cell(r, 2).fill = FILL(fill); ws.cell(r, 2).border = Border(left=fin, right=fin, top=fin, bottom=fin)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5); texte(ws, r, 3, lab); r += 1
    r += 1
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5); texte(ws, r, 2, 'Synthèse de tes choix (calculée automatiquement)', size=12, bold=True, color=NAVY2); r += 1
    for i, h in enumerate(('Onglet', 'Éléments cochés', 'Disponibles', 'Contenu'), 2):
        x = ws.cell(r, i, h); x.font = F(bold=True, color='FFFFFF'); x.fill = FILL(NAVY2)
    r += 1; r0 = r
    for inf in infos:
        q = "'" + inf['feuille'] + "'"
        ws.cell(r, 2, inf['feuille']).font = F(bold=True)
        ws.cell(r, 2).hyperlink = f"#{q}!B1"
        rng_id = f"{q}!$A${inf['debut']}:$A${inf['fin']}"; rng_oui = f"{q}!${inf['oui']}${inf['debut']}:${inf['oui']}${inf['fin']}"
        ws.cell(r, 3, f'=COUNTIFS({rng_oui},"Oui",{rng_id},"<>sec:*")').font = F()
        ws.cell(r, 4, f'=COUNTIFS({rng_id},"<>sec:*",{rng_id},"<>")').font = F()
        texte(ws, r, 5, inf['titre'], color=GRIS)
        for c in range(2, 6): ws.cell(r, c).border = BORD
        r += 1
    ws.cell(r, 2, 'Total').font = F(bold=True)
    ws.cell(r, 3, f'=SUM(C{r0}:C{r - 1})').font = F(bold=True); ws.cell(r, 4, f'=SUM(D{r0}:D{r - 1})').font = F(bold=True)
    r += 2
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5); texte(ws, r, 2, 'Journal des versions du classeur', size=12, bold=True, color=NAVY2); r += 1
    for v, d, txt in JOURNAL:
        ws.cell(r, 2, f'Version {v}').font = F(bold=True); ws.cell(r, 3, d).font = F()
        ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=5); texte(ws, r, 4, txt); ws.row_dimensions[r].height = 30; r += 1
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = '8B5E00'

def feuille_tuto(wb):
    ws = wb.create_sheet('Tuto', 1)
    for col, w in zip('ABCDE', (3, 6, 30, 28, 70)): ws.column_dimensions[col].width = w
    ws.merge_cells('B1:E1'); t = ws['B1']; t.value = 'Tuto · remplir le classeur et installer sa propre version'
    t.font = F(size=17, bold=True, color='FFFFFF'); t.alignment = MID
    for c in 'BCDE': ws[f'{c}1'].fill = FILL(NAVY)
    ws.row_dimensions[1].height = 36
    r = 3
    for titre, intro, etapes in TUTO_REMPLIR + TUTO_PRIVE:
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
        x = texte(ws, r, 2, titre, size=13, bold=True, color='FFFFFF')
        for c in 'BCDE': ws[f'{c}{r}'].fill = FILL(NAVY2)
        ws.row_dimensions[r].height = 24; r += 1
        if intro:
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5); texte(ws, r, 2, intro, italic=True, color=GRIS)
            ws.row_dimensions[r].height = max(18, 15 * (len(intro) // 115 + 1)); r += 1
        for i, h in enumerate(('N°', 'Étape', 'Où', 'Comment'), 2):
            x = ws.cell(r, i, h); x.font = F(bold=True, color=NAVY2); x.fill = FILL(SEC)
        r += 1
        for n, (etape, ou, comment, lien) in enumerate(etapes, 1):
            ws.cell(r, 2, n).font = F(bold=True, color=NAVY); ws.cell(r, 2).alignment = Alignment(horizontal='center', vertical='top')
            texte(ws, r, 3, etape, bold=True); x = texte(ws, r, 4, ou)
            if lien: x.hyperlink = lien; x.font = F(color='1F5BA8', underline='single')
            texte(ws, r, 5, comment)
            ws.row_dimensions[r].height = max(30, 14 * (len(comment) // 78 + 1) + 4)
            for c in range(2, 6): ws.cell(r, c).border = BORD
            r += 1
        r += 1
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = '8B5E00'

def construire_classeur(ref, sortie, ancien=None):
    choix, ajouts = anciens_choix(ancien)
    wb = Workbook(); wb.remove(wb.active)
    infos = [feuille_module(wb, m, choix, ajouts) for m in construire(ref)]
    feuille_lisezmoi(wb, infos); feuille_tuto(wb)
    wb.active = 0
    wb.properties.title = 'Mon Écran · configuration'; wb.properties.creator = 'Mon Écran'
    wb.save(sortie)
    return infos

if __name__ == '__main__':
    ref = json.load(open(sys.argv[1]))
    infos = construire_classeur(ref, sys.argv[2], sys.argv[3] if len(sys.argv) > 3 else None)
    print('ok', sys.argv[2], len(infos), 'onglets modules')
